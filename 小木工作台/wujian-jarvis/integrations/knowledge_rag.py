from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import sqlite3
import sys
import zipfile
from array import array
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree


DEFAULT_ROOT = Path(__file__).resolve().parents[2] / "资料包" / "资料实验基地"
DEFAULT_MODEL = "BAAI/bge-small-zh-v1.5"
TEXT_EXTENSIONS = {".md", ".txt", ".json", ".jsonl", ".csv", ".tsv", ".yaml", ".yml", ".py", ".ps1"}
TEXT_FILENAMES = {".gitignore"}
DOCUMENT_EXTENSIONS = TEXT_EXTENSIONS | {".docx", ".pdf", ".xlsx"}
CATALOG_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}
IGNORED_PARTS = {"__pycache__", ".git", ".pytest_cache", "node_modules", ".runtime"}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_text(value: str) -> str:
    value = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{4,}", "\n\n\n", value)
    return value.strip()


def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_docx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    lines = []
    for paragraph in root.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"):
        text = "".join(node.text or "" for node in paragraph.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
        if text.strip():
            lines.append(text)
    return "\n".join(lines)


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages)


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    output = []
    for sheet in book.worksheets:
        output.append(f"# 工作表：{sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value) if value is not None else "" for value in row]
            if any(values):
                output.append("\t".join(values))
    book.close()
    return "\n".join(output)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS or path.name.lower() in TEXT_FILENAMES:
        return read_text(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    raise ValueError("该文件类型仅登记，不做全文解析")


def title_from(path: Path, text: str) -> str:
    for line in text.splitlines()[:30]:
        value = re.sub(r"^#{1,6}\s*", "", line).strip()
        if value and value not in {"---", "+++"} and not value.startswith("!/") and len(value) <= 120:
            return value
    return path.stem


def repair_document_titles(connection: sqlite3.Connection) -> None:
    for row in connection.execute("SELECT id,title,path FROM documents"):
        if row["title"] in {"---", "+++"} or row["title"].startswith("!/"):
            connection.execute("UPDATE documents SET title=? WHERE id=?", (Path(row["path"]).stem, row["id"]))


def classify(relative_path: str) -> tuple[str, str, int]:
    name = relative_path.replace("\\", "/")
    lower = name.lower()
    if "历史" in name or "归档" in name or "旧" in name:
        lifecycle = "historical"
    elif "运行记录" in name:
        lifecycle = "run"
    elif "审查记录" in name or "证据" in name:
        lifecycle = "evidence"
    else:
        lifecycle = "current"

    if "系统" in name or "手册" in name or "当前" in name:
        authority, weight = "system", 100
    elif "工具" in name or "skill.md" in lower or "规则" in name or "方法" in name:
        authority, weight = "tool", 85
    elif lifecycle == "evidence":
        authority, weight = "evidence", 70
    elif "研究" in name or "报告" in name:
        authority, weight = "research", 60
    elif lifecycle == "run":
        authority, weight = "run", 45
    else:
        authority, weight = "reference", 55
    if lifecycle == "historical":
        weight = min(weight, 25)
    return lifecycle, authority, weight


def split_markdown(text: str) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    heading = "正文"
    buffer: list[str] = []
    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            if buffer:
                sections.append((heading, "\n".join(buffer).strip()))
            heading = match.group(2).strip()
            buffer = []
        else:
            buffer.append(line)
    if buffer:
        sections.append((heading, "\n".join(buffer).strip()))
    return [(heading, body) for heading, body in sections if body]


def chunk_text(text: str, suffix: str, max_chars: int = 1100, overlap: int = 140) -> list[tuple[str, str]]:
    sections = split_markdown(text) if suffix == ".md" else [("正文", text)]
    chunks: list[tuple[str, str]] = []
    for heading, body in sections:
        paragraphs = [part.strip() for part in re.split(r"\n\s*\n", body) if part.strip()]
        current = ""
        for paragraph in paragraphs:
            if len(paragraph) > max_chars:
                if current:
                    chunks.append((heading, current))
                    current = ""
                start = 0
                while start < len(paragraph):
                    end = min(len(paragraph), start + max_chars)
                    chunks.append((heading, paragraph[start:end]))
                    if end == len(paragraph):
                        break
                    start = max(0, end - overlap)
                continue
            candidate = paragraph if not current else current + "\n\n" + paragraph
            if len(candidate) > max_chars and current:
                chunks.append((heading, current))
                current = paragraph
            else:
                current = candidate
        if current:
            chunks.append((heading, current))
    return chunks


def lexical_terms(text: str) -> str:
    compact = re.sub(r"\s+", "", text.lower())
    chinese = re.findall(r"[\u3400-\u9fff]+", compact)
    grams = []
    for item in chinese:
        grams.extend(item[index:index + 2] for index in range(max(0, len(item) - 1)))
        grams.extend(item[index:index + 3] for index in range(max(0, len(item) - 2)))
    ascii_terms = re.findall(r"[a-z0-9_.%-]{2,}", text.lower())
    return " ".join(dict.fromkeys(grams + ascii_terms))


def connect(db_path: Path) -> sqlite3.Connection:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA foreign_keys=ON")
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS documents(
          id INTEGER PRIMARY KEY, path TEXT UNIQUE NOT NULL, relative_path TEXT NOT NULL,
          sha256 TEXT NOT NULL, size INTEGER NOT NULL, mtime REAL NOT NULL, suffix TEXT NOT NULL,
          title TEXT NOT NULL, lifecycle TEXT NOT NULL, authority TEXT NOT NULL,
          authority_weight INTEGER NOT NULL, parse_status TEXT NOT NULL,
          parse_error TEXT, duplicate_of INTEGER, active INTEGER NOT NULL DEFAULT 1,
          indexed_at TEXT NOT NULL, FOREIGN KEY(duplicate_of) REFERENCES documents(id)
        );
        CREATE TABLE IF NOT EXISTS chunks(
          id INTEGER PRIMARY KEY, document_id INTEGER NOT NULL, ordinal INTEGER NOT NULL,
          heading TEXT NOT NULL, content TEXT NOT NULL, lexical TEXT NOT NULL,
          embedding BLOB, embedding_dim INTEGER, embedding_model TEXT,
          FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE,
          UNIQUE(document_id, ordinal)
        );
        CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(
          chunk_id UNINDEXED, content, lexical, tokenize='unicode61'
        );
        CREATE TABLE IF NOT EXISTS index_runs(
          id INTEGER PRIMARY KEY, started_at TEXT NOT NULL, finished_at TEXT,
          mode TEXT NOT NULL, status TEXT NOT NULL, summary TEXT
        );
        CREATE TABLE IF NOT EXISTS style_profile(
          key TEXT PRIMARY KEY, value TEXT NOT NULL, source TEXT NOT NULL,
          confirmed INTEGER NOT NULL DEFAULT 1, updated_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_documents_sha ON documents(sha256);
        CREATE INDEX IF NOT EXISTS idx_documents_active ON documents(active, parse_status);
        CREATE INDEX IF NOT EXISTS idx_chunks_document ON chunks(document_id);
        """
    )
    return connection


def create_embedder(model_name: str, cache_dir: Path):
    from fastembed import TextEmbedding

    cache_dir.mkdir(parents=True, exist_ok=True)
    return TextEmbedding(model_name=model_name, cache_dir=str(cache_dir), threads=max(1, (os.cpu_count() or 2) // 2))


def pack_vector(values) -> tuple[bytes, int]:
    packed = array("f", (float(value) for value in values))
    return packed.tobytes(), len(packed)


def unpack_vector(blob: bytes) -> array:
    values = array("f")
    values.frombytes(blob)
    return values


def delete_document_chunks(connection: sqlite3.Connection, document_id: int) -> None:
    connection.execute("DELETE FROM chunks_fts WHERE chunk_id IN (SELECT id FROM chunks WHERE document_id=?)", (document_id,))
    connection.execute("DELETE FROM chunks WHERE document_id=?", (document_id,))


def seed_style_profile(connection: sqlite3.Connection) -> None:
    values = {
        "language": "默认使用简体中文",
        "answer_order": "先给结论和影响，再给必要依据",
        "futures_names": "期货品种优先使用中文名称，代码作为补充",
        "uncertainty": "资料不足时明确缺口，不猜测数据、规则或手数",
        "questions": "确需追问时一次只问一个会改变结论的关键问题",
        "approval": "模型只提供研究草案，交易批准由用户在工作台人工完成",
        "chat_import": "引用聊天01a0e190-e0cb-7072-8a41-6c14d835b36c待read_thread工具可用后导入",
    }
    now = utc_now()
    for key, value in values.items():
        connection.execute(
            "INSERT INTO style_profile(key,value,source,confirmed,updated_at) VALUES(?,?,?,?,?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value, source=excluded.source, updated_at=excluded.updated_at",
            (key, value, "当前用户确认与系统设计", 1, now),
        )


def iter_files(root: Path):
    for path in root.rglob("*"):
        if not path.is_file() or any(part in IGNORED_PARTS for part in path.parts):
            continue
        yield path


def reindex(args) -> dict:
    root = Path(args.root).resolve()
    db_path = Path(args.db).resolve()
    cache_dir = Path(args.cache).resolve()
    if not root.is_dir():
        raise ValueError(f"知识库目录不存在：{root}")
    connection = connect(db_path)
    started_at = utc_now()
    run_id = connection.execute(
        "INSERT INTO index_runs(started_at,mode,status) VALUES(?,?,?)",
        (started_at, "full" if args.full else "incremental", "running"),
    ).lastrowid
    connection.commit()
    stats = {"found": 0, "parsed": 0, "cataloged": 0, "skipped": 0, "failed": 0, "duplicates": 0, "chunks": 0, "embedded": 0, "removed": 0}
    failures = []
    seen_paths: set[str] = set()
    pending_embeddings: list[tuple[int, str]] = []
    embedder = None
    embedding_error = None
    try:
        manifest = json.loads(Path(args.manifest).read_text(encoding="utf-8")) if getattr(args, "manifest", None) else {}
        policy = {str(Path(item["path"]).resolve()).casefold(): item for item in manifest.get("files", [])}
        paths = list(dict.fromkeys(list(iter_files(root)) + [Path(item["path"]).resolve() for item in manifest.get("files", [])]))
        indexed_paths = {str(path.resolve()).casefold() for path in paths}
        for path in paths:
            stats["found"] += 1
            absolute = str(path.resolve())
            seen_paths.add(absolute.casefold())
            relative = str(path.relative_to(root)) if path.is_relative_to(root) else str(path)
            suffix = path.suffix.lower()
            can_parse = suffix in DOCUMENT_EXTENSIONS or path.name.lower() in TEXT_FILENAMES
            item_hash = sha256(path)
            item_stat = path.stat()
            existing = connection.execute("SELECT * FROM documents WHERE path=?", (absolute,)).fetchone()
            lifecycle, authority, authority_weight = classify(relative)
            if manifest:
                entry = policy.get(absolute.casefold())
                lifecycle = entry["lifecycle"] if entry else manifest.get("legacyLifecycle", "historical")
                authority = entry.get("authority", "reference") if entry else "reference"
                authority_weight = entry.get("weight", 50) if entry else 25
            if existing:
                connection.execute("UPDATE documents SET lifecycle=?,authority=?,authority_weight=? WHERE id=?", (lifecycle, authority, authority_weight, existing["id"]))
            if existing and not args.full and existing["sha256"] == item_hash and existing["active"] and existing["parse_status"] != "stale" and not (existing["parse_status"] == "cataloged" and can_parse):
                if not args.no_embeddings:
                    pending_embeddings.extend(
                        (row["id"], row["content"])
                        for row in connection.execute(
                            "SELECT id,content FROM chunks WHERE document_id=? AND embedding IS NULL",
                            (existing["id"],),
                        )
                    )
                stats["skipped"] += 1
                continue
            if existing:
                connection.execute("UPDATE documents SET duplicate_of=NULL,parse_status='stale' WHERE duplicate_of=?", (existing["id"],))
                delete_document_chunks(connection, existing["id"])
                document_id = existing["id"]
            else:
                document_id = None
            duplicate = connection.execute(
                "SELECT id,path FROM documents WHERE sha256=? AND active=1 AND path<>? ORDER BY id LIMIT 1",
                (item_hash, absolute),
            ).fetchone()
            if duplicate and duplicate["path"].casefold() not in indexed_paths:
                duplicate = None
            parse_status = "cataloged"
            parse_error = None
            text = ""
            if duplicate:
                stats["duplicates"] += 1
                parse_status = "duplicate"
            elif can_parse:
                try:
                    text = normalize_text(extract_text(path))
                    if not text:
                        parse_status = "cataloged"
                        parse_error = "未提取到可检索文字"
                        stats["cataloged"] += 1
                    else:
                        parse_status = "parsed"
                        stats["parsed"] += 1
                except Exception as error:
                    parse_status = "failed"
                    parse_error = str(error)[:500]
                    stats["failed"] += 1
                    failures.append({"path": relative, "error": parse_error})
            elif suffix in CATALOG_EXTENSIONS:
                stats["cataloged"] += 1
                parse_error = "视觉文件首版仅登记，待OCR或视觉检索"
            else:
                stats["cataloged"] += 1
                parse_error = "非首版全文解析类型，仅登记"
            title = title_from(path, text) if text else path.stem
            values = (
                absolute, relative, item_hash, item_stat.st_size, item_stat.st_mtime, suffix,
                title, lifecycle, authority, authority_weight, parse_status, parse_error,
                duplicate["id"] if duplicate else None, 1, utc_now(),
            )
            if document_id is None:
                document_id = connection.execute(
                    "INSERT INTO documents(path,relative_path,sha256,size,mtime,suffix,title,lifecycle,authority,authority_weight,parse_status,parse_error,duplicate_of,active,indexed_at) "
                    "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    values,
                ).lastrowid
            else:
                connection.execute(
                    "UPDATE documents SET path=?,relative_path=?,sha256=?,size=?,mtime=?,suffix=?,title=?,lifecycle=?,authority=?,authority_weight=?,parse_status=?,parse_error=?,duplicate_of=?,active=?,indexed_at=? WHERE id=?",
                    values + (document_id,),
                )
            if parse_status == "parsed":
                for ordinal, (heading, content) in enumerate(chunk_text(text, suffix)):
                    chunk_id = connection.execute(
                        "INSERT INTO chunks(document_id,ordinal,heading,content,lexical) VALUES(?,?,?,?,?)",
                        (document_id, ordinal, heading, content, lexical_terms(f"{title} {heading} {content}")),
                    ).lastrowid
                    connection.execute(
                        "INSERT INTO chunks_fts(chunk_id,content,lexical) SELECT id,content,lexical FROM chunks WHERE id=?",
                        (chunk_id,),
                    )
                    pending_embeddings.append((chunk_id, f"{title}\n{heading}\n{content}"))
                    stats["chunks"] += 1
            if stats["found"] % 25 == 0:
                connection.commit()
        active_rows = connection.execute("SELECT id,path FROM documents WHERE active=1").fetchall()
        for row in active_rows:
            if row["path"].casefold() not in seen_paths:
                connection.execute("UPDATE documents SET active=0,lifecycle='historical',authority_weight=25 WHERE id=?", (row["id"],))
                stats["removed"] += 1
        if pending_embeddings and not args.no_embeddings:
            try:
                embedder = create_embedder(args.model, cache_dir)
            except Exception as error:
                embedding_error = str(error)[:500]
        if embedder and pending_embeddings:
            texts = [text for _, text in pending_embeddings]
            for (chunk_id, _), vector in zip(pending_embeddings, embedder.embed(texts, batch_size=32)):
                blob, dimension = pack_vector(vector)
                connection.execute(
                    "UPDATE chunks SET embedding=?,embedding_dim=?,embedding_model=? WHERE id=?",
                    (blob, dimension, args.model, chunk_id),
                )
                stats["embedded"] += 1
                if stats["embedded"] % 100 == 0:
                    connection.commit()
        repair_document_titles(connection)
        seed_style_profile(connection)
        missing_embeddings = connection.execute("SELECT COUNT(*) FROM chunks WHERE embedding IS NULL").fetchone()[0]
        embedding_status = "ready" if not missing_embeddings else "degraded"
        connection.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('root',?)", (str(root),))
        connection.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('model',?)", (args.model,))
        connection.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('last_indexed_at',?)", (utc_now(),))
        connection.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('embedding_status',?)", (embedding_status,))
        connection.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('embedding_error',?)", (embedding_error or "",))
        summary = {**stats, "embeddingStatus": embedding_status, "embeddingError": embedding_error, "failures": failures[:50]}
        connection.execute(
            "UPDATE index_runs SET finished_at=?,status='completed',summary=? WHERE id=?",
            (utc_now(), json.dumps(summary, ensure_ascii=False), run_id),
        )
        connection.commit()
        return summary
    except Exception as error:
        connection.execute(
            "UPDATE index_runs SET finished_at=?,status='failed',summary=? WHERE id=?",
            (utc_now(), json.dumps({"error": str(error)}, ensure_ascii=False), run_id),
        )
        connection.commit()
        raise
    finally:
        connection.close()


def cosine(left: array, right: array) -> float:
    if len(left) != len(right) or not left:
        return 0.0
    return sum(a * b for a, b in zip(left, right))


def status(args) -> dict:
    db_path = Path(args.db).resolve()
    if not db_path.exists():
        return {"ready": False, "state": "not-built", "database": str(db_path)}
    connection = connect(db_path)
    counts = connection.execute(
        "SELECT COUNT(*) total, SUM(active) active, SUM(active AND parse_status='parsed') parsed, "
        "SUM(active AND parse_status='cataloged') cataloged, SUM(active AND parse_status='failed') failed, "
        "SUM(active AND parse_status='duplicate') duplicates FROM documents"
    ).fetchone()
    chunks = connection.execute("SELECT COUNT(*) count, SUM(embedding IS NOT NULL) embedded FROM chunks").fetchone()
    meta = {row["key"]: row["value"] for row in connection.execute("SELECT key,value FROM meta")}
    latest = connection.execute("SELECT * FROM index_runs ORDER BY id DESC LIMIT 1").fetchone()
    result = {
        "ready": bool(counts["parsed"] and chunks["count"]),
        "state": "ready" if counts["parsed"] and chunks["count"] else "empty",
        "database": str(db_path),
        "root": meta.get("root", str(DEFAULT_ROOT)),
        "model": meta.get("model", DEFAULT_MODEL),
        "embeddingStatus": meta.get("embedding_status", "unknown"),
        "embeddingError": meta.get("embedding_error") or None,
        "lastIndexedAt": meta.get("last_indexed_at"),
        "documents": {key: int(counts[key] or 0) for key in ("total", "active", "parsed", "cataloged", "failed", "duplicates")},
        "chunks": int(chunks["count"] or 0),
        "embeddedChunks": int(chunks["embedded"] or 0),
        "latestRun": dict(latest) if latest else None,
    }
    connection.close()
    return result


def search(args) -> dict:
    query = normalize_text(args.query)
    if not query or len(query) > 500:
        raise ValueError("检索问题为空或过长")
    retrieval_query = re.sub(r"根据(?:我的|本地|当前)?(?:规则|资料|知识库)|请问|帮我|告诉我|是什么|有哪些|如何|怎么|[？?]", "", query).strip()
    if len(retrieval_query) < 2:
        retrieval_query = query
    db_path = Path(args.db).resolve()
    if not db_path.exists():
        return {"query": query, "hits": [], "degraded": True, "notice": "知识库尚未建立"}
    connection = connect(db_path)
    meta = {row["key"]: row["value"] for row in connection.execute("SELECT key,value FROM meta")}
    candidates: dict[int, dict] = {}
    query_lexical = lexical_terms(retrieval_query)
    tokens = query_lexical.split()
    if tokens:
        expression = " OR ".join(f'"{token.replace(chr(34), "")}"' for token in tokens[:24])
        try:
            rows = connection.execute(
                "SELECT chunk_id,bm25(chunks_fts) score FROM chunks_fts WHERE chunks_fts MATCH ? ORDER BY score LIMIT 80",
                (expression,),
            ).fetchall()
            for rank, row in enumerate(rows):
                candidates.setdefault(int(row["chunk_id"]), {})["ftsRank"] = 1.0 / (rank + 1)
        except sqlite3.OperationalError:
            pass
    like_rows = connection.execute(
        "SELECT c.id FROM chunks c JOIN documents d ON d.id=c.document_id WHERE d.active=1 AND (c.content LIKE ? OR d.title LIKE ? OR d.relative_path LIKE ?) LIMIT 80",
        (f"%{retrieval_query}%", f"%{retrieval_query}%", f"%{retrieval_query}%"),
    ).fetchall()
    for rank, row in enumerate(like_rows):
        candidates.setdefault(int(row["id"]), {})["exact"] = 1.0 / (rank + 1)
    degraded = True
    embedding_error = None
    vector_rows = connection.execute("SELECT c.id,c.embedding FROM chunks c JOIN documents d ON d.id=c.document_id WHERE d.active=1 AND c.embedding IS NOT NULL").fetchall()
    if vector_rows:
        try:
            embedder = create_embedder(meta.get("model", DEFAULT_MODEL), Path(args.cache).resolve())
            query_vector = next(iter(embedder.query_embed(retrieval_query)))
            packed_query = array("f", (float(value) for value in query_vector))
            scored = sorted(((cosine(packed_query, unpack_vector(row["embedding"])), int(row["id"])) for row in vector_rows), reverse=True)[:80]
            for rank, (score, chunk_id) in enumerate(scored):
                if score > 0:
                    candidates.setdefault(chunk_id, {})["semantic"] = score
                    candidates[chunk_id]["semanticRank"] = 1.0 / (rank + 1)
            degraded = False
        except Exception as error:
            embedding_error = str(error)[:300]
    if not candidates:
        connection.close()
        return {"query": query, "hits": [], "degraded": degraded, "notice": "本地知识库没有足够证据", "embeddingError": embedding_error}
    ids = list(candidates)
    placeholders = ",".join("?" for _ in ids)
    rows = connection.execute(
        f"SELECT c.id,c.heading,c.content,c.lexical,d.id document_id,d.path,d.relative_path,d.title,d.sha256,d.lifecycle,d.authority,d.authority_weight,d.indexed_at "
        f"FROM chunks c JOIN documents d ON d.id=c.document_id WHERE d.active=1 AND c.id IN ({placeholders})",
        ids,
    ).fetchall()
    ranked = []
    for row in rows:
        signals = candidates[int(row["id"])]
        semantic = float(signals.get("semantic", 0.0))
        query_terms = set(tokens)
        chunk_terms = set(row["lexical"].split()) | set(lexical_terms(row["relative_path"]).split())
        lexical = sum(2 if len(term) >= 3 else 1 for term in query_terms & chunk_terms) / max(1, sum(2 if len(term) >= 3 else 1 for term in query_terms))
        exact = float(signals.get("exact", 0.0))
        authority = int(row["authority_weight"]) / 100
        current = 1.0 if row["lifecycle"] == "current" else 0.0
        score = semantic * 0.62 + float(signals.get("semanticRank", 0.0)) * 0.03 + lexical * 0.20 + exact * 0.02 + authority * 0.08 + current * 0.05
        if not semantic:
            score = lexical * 0.68 + exact * 0.12 + authority * 0.12 + current * 0.08
        sources = [source["path"] for source in connection.execute("SELECT path FROM documents WHERE sha256=? AND active=1 ORDER BY path", (row["sha256"],))]
        ranked.append({
            "chunkId": int(row["id"]), "documentId": int(row["document_id"]), "title": row["title"],
            "heading": row["heading"], "path": row["path"], "relativePath": row["relative_path"],
            "sourcePaths": sources, "lifecycle": row["lifecycle"], "authority": row["authority"],
            "indexedAt": row["indexed_at"], "score": round(score, 6), "semanticScore": round(semantic, 6),
            "excerpt": row["content"][:1000],
        })
    ranked.sort(key=lambda item: item["score"], reverse=True)
    selected = []
    per_document: dict[int, int] = {}
    for item in ranked:
        if per_document.get(item["documentId"], 0) >= 2:
            continue
        if item["score"] < (0.18 if degraded else 0.26):
            continue
        selected.append(item)
        per_document[item["documentId"]] = per_document.get(item["documentId"], 0) + 1
        if len(selected) >= args.limit:
            break
    connection.close()
    return {
        "query": query, "hits": selected, "degraded": degraded, "embeddingError": embedding_error,
        "notice": None if selected else "本地知识库没有足够证据",
    }


def source(args) -> dict:
    connection = connect(Path(args.db).resolve())
    row = connection.execute(
        "SELECT c.id,c.heading,c.content,d.path,d.relative_path,d.title,d.lifecycle,d.authority,d.indexed_at "
        "FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.id=? AND d.active=1",
        (args.chunk_id,),
    ).fetchone()
    connection.close()
    if not row:
        raise ValueError("知识片段不存在或已经失效")
    return dict(row)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--db", required=True)
    parser.add_argument("--cache", required=True)
    parser.add_argument("--root", default=str(DEFAULT_ROOT))
    parser.add_argument("--model", default=DEFAULT_MODEL)
    parser.add_argument("--manifest")
    subparsers = parser.add_subparsers(dest="command", required=True)
    index_parser = subparsers.add_parser("reindex")
    index_parser.add_argument("--full", action="store_true")
    index_parser.add_argument("--no-embeddings", action="store_true")
    search_parser = subparsers.add_parser("search")
    search_parser.add_argument("query")
    search_parser.add_argument("--limit", type=int, default=6)
    source_parser = subparsers.add_parser("source")
    source_parser.add_argument("chunk_id", type=int)
    subparsers.add_parser("status")
    args = parser.parse_args()
    if args.command == "reindex":
        result = reindex(args)
    elif args.command == "search":
        args.limit = max(1, min(args.limit, 12))
        result = search(args)
    elif args.command == "source":
        result = source(args)
    else:
        result = status(args)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": str(error)}, ensure_ascii=False))
        sys.exit(1)
